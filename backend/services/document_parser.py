"""
文档解析器 - 支持 PDF、TXT、Markdown、DOCX、网页 URL
"""

import PyPDF2
import requests
from typing import List
from langchain.text_splitter import RecursiveCharacterTextSplitter
from docx import Document
from bs4 import BeautifulSoup

from core.config import settings


class DocumentParser:
    """文档解析器"""

    def __init__(self):
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.CHUNK_SIZE,
            chunk_overlap=settings.CHUNK_OVERLAP,
            separators=["\n\n", "\n", "。", ".", " ", ""]
        )

    def parse_file(self, file_path: str) -> str:
        """
        解析文件，提取文本内容

        Args:
            file_path: 文件路径

        Returns:
            text: 提取的文本
        """
        file_ext = file_path.split(".")[-1].lower()

        if file_ext == "pdf":
            return self._parse_pdf(file_path)
        elif file_ext in ["txt", "md"]:
            return self._parse_text(file_path)
        elif file_ext == "docx":
            return self._parse_docx(file_path)
        else:
            raise ValueError(f"不支持的文件格式: {file_ext}")

    def parse_url(self, url: str) -> str:
        """
        解析网页 URL，提取文本内容

        Args:
            url: 网页 URL

        Returns:
            text: 提取的文本
        """
        try:
            # 发送 HTTP 请求
            headers = {
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'
            }
            response = requests.get(url, headers=headers, timeout=10)
            response.raise_for_status()

            # 使用 BeautifulSoup 解析 HTML
            soup = BeautifulSoup(response.content, 'html.parser')

            # 移除 script 和 style 标签
            for script in soup(["script", "style"]):
                script.decompose()

            # 提取文本
            text = soup.get_text(separator='\n', strip=True)

            # 添加 URL 信息
            text = f"[Source URL: {url}]\n\n{text}"

            return text

        except Exception as e:
            raise ValueError(f"Failed to fetch URL: {str(e)}")

    def _parse_pdf(self, file_path: str) -> str:
        """
        解析 PDF 文件

        Args:
            file_path: PDF 文件路径

        Returns:
            text: 提取的文本
        """
        text = ""
        with open(file_path, "rb") as file:
            pdf_reader = PyPDF2.PdfReader(file)
            for page_num, page in enumerate(pdf_reader.pages):
                page_text = page.extract_text()
                # 保留页码信息
                text += f"\n[Page {page_num + 1}]\n{page_text}\n"

        return text

    def _parse_text(self, file_path: str) -> str:
        """
        解析文本文件

        Args:
            file_path: 文本文件路径

        Returns:
            text: 文本内容
        """
        with open(file_path, "r", encoding="utf-8") as file:
            return file.read()

    def _parse_docx(self, file_path: str) -> str:
        """
        解析 DOCX 文件

        Args:
            file_path: DOCX 文件路径

        Returns:
            text: 提取的文本
        """
        doc = Document(file_path)

        # 提取段落文本
        text = ""
        for para in doc.paragraphs:
            if para.text.strip():
                text += para.text + "\n"

        # 提取表格文本
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells])
                text += row_text + "\n"

        return text

    def chunk_text(self, text: str) -> List[str]:
        """
        将文本分块

        Args:
            text: 原始文本

        Returns:
            chunks: 文本块列表
        """
        chunks = self.text_splitter.split_text(text)
        return chunks
